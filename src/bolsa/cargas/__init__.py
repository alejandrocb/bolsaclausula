"""Cargadores de ficheros (iniciales inmutables y periódicos)."""
from __future__ import annotations

import csv
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Iterator


def lee_csv(ruta, delimitador: str) -> list[dict]:
    """Lee un CSV a lista de dicts, tolerante a BOM."""
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=delimitador))


_NULOS = {"", "null", "[null]", "(null)", "none", "nan", "\\n"}


def limpio(valor) -> str:
    """Normaliza texto tratando marcadores de NULL como cadena vacía."""
    if valor is None:
        return ""
    txt = str(valor).strip()
    return "" if txt.lower() in _NULOS else txt


def entero(valor, defecto: int = 0) -> int:
    """Convierte a int de forma tolerante ('', None, '  12 ' -> 12)."""
    if valor is None:
        return defecto
    txt = str(valor).strip()
    if txt == "":
        return defecto
    try:
        return int(float(txt.replace(",", ".")))
    except ValueError:
        return defecto


# --- ODS (OpenDocument Spreadsheet) sin dependencias ---------------------------
_NS_T = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
_NS_TX = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"


def _texto_celda(celda) -> str:
    return " ".join("".join(p.itertext()) for p in celda.iter(f"{{{_NS_TX}}}p"))


def lee_ods(ruta, hoja: str | None = None) -> list[list[str]]:
    """Devuelve las filas (lista de celdas de texto) de una hoja ODS."""
    with zipfile.ZipFile(ruta) as z:
        root = ET.fromstring(z.read("content.xml").decode("utf-8"))
    tablas = list(root.iter(f"{{{_NS_T}}}table"))
    tabla = None
    for t in tablas:
        if hoja is None or t.get(f"{{{_NS_T}}}name") == hoja:
            tabla = t
            break
    # si la hoja pedida no existe, usamos la primera (el export puede
    # nombrarla distinto), sin fallar.
    if tabla is None and tablas:
        tabla = tablas[0]
    if tabla is None:
        raise ValueError(f"No hay hojas en {ruta}")
    filas: list[list[str]] = []
    for fila in tabla.iter(f"{{{_NS_T}}}table-row"):
        celdas: list[str] = []
        for c in fila.findall(f"{{{_NS_T}}}table-cell"):
            rep = int(c.get(f"{{{_NS_T}}}number-columns-repeated", "1"))
            txt = _texto_celda(c)
            celdas.extend([txt] * rep)
        while celdas and celdas[-1] == "":
            celdas.pop()
        if celdas:
            filas.append(celdas)
    return filas


def lee_xlsx(ruta, hoja: str | None = None) -> list[list]:
    """Filas (valores crudos) de una hoja XLSX. Acepta cualquier extensión
    (lee el contenido en memoria, sin depender del sufijo del fichero)."""
    import io
    import openpyxl

    wb = openpyxl.load_workbook(
        io.BytesIO(Path(ruta).read_bytes()), read_only=True, data_only=True
    )
    # si la hoja pedida no existe (el export puede llamarla distinto:
    # 'Query', 'Export', 'Hoja1'…) usamos la primera, sin fallar.
    ws = wb[hoja] if (hoja and hoja in wb.sheetnames) else wb[wb.sheetnames[0]]
    filas = []
    for fila in ws.iter_rows(values_only=True):
        if fila is None or all(c is None for c in fila):
            continue
        filas.append(list(fila))
    return filas


def lee_tabla(ruta, hoja: str | None = None) -> list[list]:
    """Lee una hoja de cálculo autodetectando XLSX vs ODS por su contenido
    (algunos export vienen con extensión .ods pero son XLSX y viceversa)."""
    with zipfile.ZipFile(ruta) as z:
        nombres = z.namelist()
    if "content.xml" in nombres:
        return lee_ods(ruta, hoja)
    return lee_xlsx(ruta, hoja)


def dicts_desde_filas(filas: list[list[str]]) -> Iterator[dict]:
    """Convierte filas [cabecera, *datos] en dicts."""
    if not filas:
        return
    cabecera = filas[0]
    for fila in filas[1:]:
        fila = fila + [""] * (len(cabecera) - len(fila))
        yield dict(zip(cabecera, fila))
