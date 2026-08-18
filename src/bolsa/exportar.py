"""Exportación a XLSX: cuaderno de conciliación + fichero de recarga."""
from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import CONFIG
from .equivalencias import Equivalencias
from .motor import ResultadoConciliacion

_CAB = Font(bold=True, color="FFFFFF")
_FILL_CAB = PatternFill("solid", fgColor="305496")
_COLOR = {
    "ROJO": PatternFill("solid", fgColor="FFC7CE"),
    "AMBAR": PatternFill("solid", fgColor="FFEB9C"),
    "VERDE": PatternFill("solid", fgColor="C6EFCE"),
    "SIN DATOS": PatternFill("solid", fgColor="D9D9D9"),
}

# Columnas del saldo inicial 01/01/2026 reconstruido (anclado a PeopleNet)
_COLS_SALDO_INI_DIR = [
    ("clausula", "Cláusula"), ("direccion", "Dir."),
    ("direccion_nombre", "Dirección"),
    ("saldo_01_01", "Saldo inicial 01/01/2026"),
    ("usados_corte", "Consumo PeopleNet al corte (02/07)"),
    ("saldo_corte", "Saldo en el corte (02/07)"),
    ("usados_desde", "Consumo PeopleNet desde el corte"),
    ("saldo_actual", "Saldo actual (solo PeopleNet)"),
]
_COLS_SALDO_INI_PLAZA = [
    ("clausula", "Cláusula"), ("direccion", "Dir."),
    ("direccion_nombre", "Dirección"), ("id_plaza", "ID Plaza"),
    ("saldo_01_01", "Saldo inicial 01/01/2026"),
    ("usados_corte", "Consumo PeopleNet al corte (02/07)"),
    ("saldo_corte", "Saldo en el corte (02/07)"),
    ("usados_desde", "Consumo PeopleNet desde el corte"),
    ("saldo_actual", "Saldo actual (solo PeopleNet)"),
]


def _hoja(wb: Workbook, titulo: str, columnas: list[tuple[str, str]], filas: list[dict]):
    ws = wb.create_sheet(titulo[:31])
    for j, (_, tit) in enumerate(columnas, start=1):
        c = ws.cell(row=1, column=j, value=tit)
        c.font = _CAB
        c.fill = _FILL_CAB
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, fila in enumerate(filas, start=2):
        for j, (clave, _) in enumerate(columnas, start=1):
            ws.cell(row=i, column=j, value=fila.get(clave))
    ws.freeze_panes = "A2"
    # anchos aproximados
    for j, (clave, tit) in enumerate(columnas, start=1):
        ancho = max(len(tit), 10)
        for fila in filas[:200]:
            ancho = max(ancho, len(str(fila.get(clave, ""))))
        ws.column_dimensions[get_column_letter(j)].width = min(ancho + 2, 45)
    return ws


def _fd(f) -> dict:
    """FilaDetalle -> dict."""
    return dict(
        id_plaza=f.id_plaza, direccion_codigo=f.direccion_codigo,
        direccion_nombre=f.direccion_nombre, descripcion_categoria=f.descripcion_categoria,
        clausula=f.clausula, saldo_base=f.saldo_base,
        control_inicial_neto=f.control_inicial_neto, saldo_postcontrol=f.saldo_postcontrol,
        consumo_posterior=f.consumo_posterior, devolucion_posterior=f.devolucion_posterior,
        devolucion_cierre=f.devolucion_cierre,
        ajuste_peoplenet=f.ajuste_peoplenet, reserva_pendiente=f.reserva_pendiente,
        saldo_calculado=f.saldo_calculado, saldo_actual_propuestas=f.saldo_actual_propuestas,
        diferencia=f.diferencia, nota=f.nota,
    )


def exporta_conciliacion(
    res: ResultadoConciliacion,
    ruta: str | Path,
    saldo_inicial: Optional[dict] = None,
    equivalencias: Optional[Equivalencias] = None,
    importaciones: Optional[list[dict]] = None,
    config=CONFIG,
) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    # 1) Detalle
    _hoja(wb, "1_Detalle", [
        ("id_plaza", "ID Plaza"), ("descripcion_categoria", "Categoría"),
        ("direccion_codigo", "Dir."), ("direccion_nombre", "Dirección"),
        ("clausula", "Cláusula"), ("saldo_base", "Saldo base 02/07"),
        ("control_inicial_neto", "Control inicial (neto)"),
        ("saldo_postcontrol", "Saldo postcontrol"),
        ("consumo_posterior", "Consumo posterior"),
        ("devolucion_posterior", "Devolución posterior"),
        ("devolucion_cierre", "Devolución por cierre (calc.)"),
        ("ajuste_peoplenet", "Ajuste PeopleNet (cláu/dir)"),
        ("reserva_pendiente", "Reserva pendiente"),
        ("saldo_calculado", "Saldo calculado (recargar)"),
        ("saldo_actual_propuestas", "Saldo actual Propuestas"),
        ("diferencia", "Diferencia"), ("nota", "Nota"),
    ], [_fd(f) for f in res.detalle])

    # 2) Resumen Dirección + cláusula
    _hoja(wb, "2_Resumen_Direccion", [
        ("direccion_codigo", "Dir."), ("direccion_nombre", "Dirección"),
        ("clausula", "Cláusula"), ("saldo_postcontrol", "Saldo postcontrol"),
        ("consumo_posterior", "Consumo"), ("devolucion_posterior", "Devolución"),
        ("devolucion_cierre", "Devol. cierre"),
        ("ajuste_peoplenet", "Ajuste PeopleNet"), ("reserva_pendiente", "Reserva pend."),
        ("saldo_calculado", "Saldo calculado"),
        ("saldo_actual_propuestas", "Saldo Propuestas"), ("diferencia", "Diferencia"),
    ], res.resumen_direccion)

    # 3) Totales por cláusula
    _hoja(wb, "3_Totales_Clausula", [
        ("clausula", "Cláusula"), ("saldo_postcontrol", "Saldo postcontrol"),
        ("consumo_posterior", "Consumo"), ("devolucion_posterior", "Devolución"),
        ("devolucion_cierre", "Devol. cierre"),
        ("ajuste_peoplenet", "Ajuste PeopleNet"), ("reserva_pendiente", "Reserva pend."),
        ("saldo_calculado", "Saldo calculado"),
        ("saldo_actual_propuestas", "Saldo Propuestas"),
    ], res.totales_clausula)

    # 4) Semáforo
    cols_sem = [
        ("clausula", "Cláusula"),
        ("bolsa_oficial_contratacion", "PeopleNet contratación"),
        ("bolsa_oficial_usados", "PeopleNet usados (contratos)"),
        ("disponible_peoplenet", "Disponible PeopleNet"),
        ("reserva_pendiente_sin_contrato", "Reservas sin contrato"),
        ("comprometido_total", "Comprometido total (usados+reservas)"),
        ("disponible_tras_compromisos", "Margen tras compromisos"),
        ("saldo_calculado_recarga", "Saldo recarga (por plaza)"),
        ("saldo_actual_propuestas", "Saldo actual Propuestas"),
        ("estado", "ESTADO"),
    ]
    ws = _hoja(wb, "4_Semaforo", cols_sem, res.semaforo)
    col_estado = len(cols_sem)
    for i, s in enumerate(res.semaforo, start=2):
        ws.cell(row=i, column=col_estado).fill = _COLOR.get(s["estado"], _COLOR["SIN DATOS"])
        ws.cell(row=i, column=col_estado).font = Font(bold=True)

    # --- Auditoría ---
    dps = res.detalle_propuestas

    def dp(d):
        return dict(
            propuesta_id=d.propuesta_id, idrh=d.idrh, id_plaza=d.id_plaza,
            direccion_codigo=d.direccion_codigo,
            clausula_declarada=d.clausula_declarada, clausula_efectiva=d.clausula_efectiva,
            fecha_inicio=str(d.fecha_inicio or ""), reserva_fin=str(d.reserva_fin or ""),
            efectivo_fin=str(d.efectivo_fin or ""), consumo_bruto=d.consumo_bruto,
            consumo_neto=d.consumo_neto, devolucion_prevista=d.devolucion_prevista,
            devolucion_registrada=d.devolucion_registrada,
            devolucion_pendiente=d.devolucion_pendiente,
            movimiento_importe=d.movimiento_importe, mecanizada=d.mecanizada,
            sub_estado=d.sub_estado,
            enlazada=d.enlazada, contrato_cerrado=d.contrato_cerrado,
            enlace_directo=d.enlace_directo,
            direccion_declarada=d.direccion_declarada,
            direccion_efectiva=d.direccion_efectiva,
            dias_sin_tramo=d.dias_sin_tramo,
            motivo_no_computa=d.motivo_no_computa,
        )

    cols_prop = [
        ("propuesta_id", "Propuesta"), ("idrh", "IDRH"), ("id_plaza", "ID Plaza"),
        ("sub_estado", "Sub_estado"), ("enlace_directo", "Enlace directo"),
        ("direccion_declarada", "Dir. prop."), ("direccion_efectiva", "Dir. real"),
        ("clausula_declarada", "Cláusula prop."),
        ("clausula_efectiva", "Cláusula real"), ("fecha_inicio", "Inicio"),
        ("reserva_fin", "Fin reservado"), ("efectivo_fin", "Fin computable"),
        ("consumo_bruto", "Consumo bruto"), ("consumo_neto", "Consumo neto"),
        ("devolucion_prevista", "Devol. prevista"),
        ("devolucion_registrada", "Devol. registrada"),
        ("devolucion_pendiente", "Devol. pendiente"),
        ("dias_sin_tramo", "Días sin tramo GFH"),
        ("movimiento_importe", "Σ movimientos"),
    ]

    reservas = [dp(d) for d in dps if d.computa and not d.mecanizada]
    _hoja(wb, "A1_Reservas_pdte_mecanizar", cols_prop, reservas)

    mec_sin = [dp(d) for d in dps if d.computa and d.mecanizada and not d.enlazada]
    _hoja(wb, "A2_Mecanizadas_sin_contrato", cols_prop, mec_sin)

    # Resumen de la reserva pendiente (sin contrato) por sub_estado
    from collections import defaultdict as _dd
    resu = _dd(lambda: {"n": 0, "dias": 0})
    for d in dps:
        if d.computa and not d.enlazada:
            k = (d.clausula_efectiva, d.sub_estado or "(vacío)")
            resu[k]["n"] += 1
            resu[k]["dias"] += d.consumo_neto
    _hoja(wb, "A2b_Reserva_por_subestado", [
        ("clausula", "Cláusula"), ("sub_estado", "Sub_estado"),
        ("n", "Nº propuestas"), ("dias", "Días netos"),
    ], [dict(clausula=k[0], sub_estado=k[1], n=v["n"], dias=v["dias"])
        for k, v in sorted(resu.items())])

    distinta = [dp(d) for d in dps if d.computa and d.clausula_distinta]
    _hoja(wb, "A3_Clausula_distinta", cols_prop, distinta)

    dir_dist = [dp(d) for d in dps
                if d.computa and (d.direccion_distinta or d.dias_sin_tramo > 0)]
    _hoja(wb, "A3b_Direccion_distinta", cols_prop, dir_dist)

    _hoja(wb, "A4_Contratos_sin_propuesta", [
        ("idrh", "IDRH"), ("num_periodo", "Núm. periodo"), ("id_plaza", "ID Plaza"),
        ("clausula", "Cláusula"), ("fecha_inicio", "Inicio"), ("fecha_fin", "Fin"),
        ("motivo_inicio", "Motivo"),
    ], [dict(idrh=c.idrh, num_periodo=c.num_periodo, id_plaza=c.id_plaza,
            clausula=c.clausula, fecha_inicio=str(c.fecha_inicio or ""),
            fecha_fin=str(c.fecha_fin or ""), motivo_inicio=c.motivo_inicio)
        for c in res.contratos_sin_propuesta])

    cierres = [dp(d) for d in dps if d.computa and d.contrato_cerrado
               and d.devolucion_prevista > 0]
    _hoja(wb, "A5_Cierres_devoluciones", cols_prop, cierres)

    _hoja(wb, "A4b_Contratos_post_corte", [
        ("idrh", "IDRH"), ("num_periodo", "Nº periodo"), ("id_plaza", "ID Plaza"),
        ("clausula", "Cláusula"), ("inicio_plaza", "Inicio plaza"),
        ("fin_plaza", "Fin plaza"), ("alta", "Alta (congelada)"),
        ("primera_aparicion", "1ª aparición (importación)"),
        ("nuevo_en_importacion", "¿Nuevo esta importación?"),
        ("propuesta_ref", "Propuesta (comentario)"),
        ("alta_posterior_corte", "¿Alta ≥ 02/07?"),
        ("enlazado_a_propuesta", "¿Enlaza propuesta viva?"),
        ("incidencia", "Incidencia"),
    ], res.contratos_post_corte)

    _hoja(wb, "A5b_Devoluciones_cierre", [
        ("propuesta_id", "Propuesta"), ("idrh", "IDRH"), ("id_plaza", "ID Plaza"),
        ("direccion", "Dirección"), ("clausula", "Cláusula"),
        ("contrato_periodo", "Nº periodo"), ("contrato_fin", "Cierre contrato"),
        ("reservado_hasta", "Reservado hasta"), ("dias_devueltos", "Días devueltos (calc.)"),
        ("devolucion_registrada", "Devol. registrada (mov.)"),
        ("devolucion_pendiente", "Devol. pendiente"),
    ], res.devoluciones_cierre)

    con_mov = [dp(d) for d in dps if d.movimiento_importe != 0
               or d.devolucion_registrada != 0]
    _hoja(wb, "A6_Movimientos_por_propuesta", cols_prop, con_mov)

    # A7 Auditoría saldo inicial y postcontrol
    if saldo_inicial:
        aud = []
        for si in saldo_inicial.values():
            if si.clave.clausula not in config.clausulas_prioritarias:
                continue
            aud.append(dict(
                id_plaza=si.clave.id_plaza, direccion_codigo=si.clave.direccion_codigo,
                clausula=si.clave.clausula, saldo_base_02_07=si.saldo_base_02_07,
                nuevos_control=si.nuevos_control, recuperados_control=si.recuperados_control,
                saldo_postcontrol=si.saldo_postcontrol, cuadra=si.cuadra(),
                origen_ajustes=si.origen_ajustes,
            ))
        _hoja(wb, "A7_Auditoria_corte", [
            ("id_plaza", "ID Plaza"), ("direccion_codigo", "Dir."),
            ("clausula", "Cláusula"), ("saldo_base_02_07", "Saldo base 02/07"),
            ("nuevos_control", "Nuevos control"),
            ("recuperados_control", "Recuperados control"),
            ("saldo_postcontrol", "Saldo postcontrol"), ("cuadra", "¿Cuadra?"),
            ("origen_ajustes", "Origen ajustes"),
        ], aud)

    # A8 Histórico de importaciones + equivalencias
    _hoja(wb, "A8_Importaciones", [
        ("fichero", "Fichero"), ("tipo", "Tipo"), ("filas", "Filas"),
        ("hash", "Hash SHA-256"), ("fecha_importacion", "Importado"),
    ], importaciones or [])
    if equivalencias:
        _hoja(wb, "A8b_Equivalencias_NIE_DNI", [
            ("idrh_a", "IDRH A"), ("idrh_b", "IDRH B"), ("motivo", "Motivo"),
            ("fecha_alta", "Alta"), ("usuario", "Usuario"),
        ], equivalencias.registro)

    # A9 Excepciones que requieren comprobación
    exc = []
    for d in dps:
        problemas = []
        if d.laboral and d.motivo_no_computa.startswith("plaza laboral") and d.consumo_bruto:
            problemas.append("propuesta sobre plaza laboral")
        if d.computa and not d.idrh:
            problemas.append("sin IDRH (no enlazable)")
        if d.computa and not d.enlazada:
            problemas.append("sin contrato localizado")
        if d.computa and d.clausula_distinta:
            problemas.append("cláusula distinta")
        if d.computa and d.plaza_distinta:
            problemas.append("enlace por plaza equivalente")
        if d.computa and d.direccion_distinta:
            problemas.append("dirección distinta (contrato manda)")
        if d.computa and d.dias_sin_tramo > 0:
            problemas.append("días sin tramo GFH que los cubra")
        if d.computa and d.devolucion_pendiente > 0:
            problemas.append("devolución prevista no registrada")
        if problemas:
            row = dp(d)
            row["incidencias"] = "; ".join(problemas)
            exc.append(row)
    _hoja(wb, "A9_Excepciones", cols_prop + [("incidencias", "Incidencias")], exc)

    # A10 Validación esperado vs registrado (por propuesta)
    val = res.validacion
    if val is not None:
        _hoja(wb, "A10_Esperado_vs_Registrado", [
            ("propuesta_id", "Propuesta"), ("idrh", "IDRH"), ("id_plaza", "ID Plaza"),
            ("direccion_codigo", "Dir."), ("clausula", "Cláusula"),
            ("computa", "¿Computa?"), ("esperado_con_tope", "Esperado (tope 31/12)"),
            ("esperado_sin_tope", "Esperado (sin tope)"),
            ("registrado_nueva", "Registrado (NUEVA)"), ("diferencia", "Diferencia"),
            ("devolucion_prevista", "Devol. prevista"),
            ("devolucion_registrada", "Devol. registrada"),
            ("devolucion_pendiente", "Devol. pendiente"), ("nota", "Nota"),
        ], [vars(f) for f in val.por_propuesta])

        # A11 Movimientos anómalos
        _hoja(wb, "A11_Movimientos_anomalos", [
            ("movimiento_id", "Movimiento"), ("propuesta_id", "Propuesta"),
            ("tipo_movimiento", "Tipo"), ("importe", "Importe"),
            ("id_plaza", "ID Plaza"), ("direccion_codigo", "Dir."),
            ("clausula", "Cláusula"), ("incidencias", "Incidencias"),
        ], [vars(a) for a in val.anomalos])

    # A12 Coherencia del enlace directo (por comentario del contrato):
    # ¿el contrato referenciado corresponde en DNI, fecha de inicio y plaza?
    _ETIQ_COH = {
        "distinto": "DISTINTO", "nie_dni": "NIE↔DNI (posible mismo)",
        "laboral_estatutario": "L↔E (misma categoría)",
    }
    incoh = []
    for d in res.detalle_propuestas:
        if not d.enlace_directo or d.coherencia_ok:
            continue
        probl = []
        if d.coh_dni not in ("ok", "sin_dni"):
            probl.append(f"DNI {_ETIQ_COH.get(d.coh_dni, d.coh_dni)}")
        if d.coh_fecha != "ok":
            probl.append("FECHA INICIO distinta")
        if d.coh_plaza != "ok":
            probl.append(f"PLAZA {_ETIQ_COH.get(d.coh_plaza, d.coh_plaza)}")
        incoh.append(dict(
            propuesta_id=d.propuesta_id,
            idrh_prop=d.idrh, idrh_contrato=d.contrato_idrh,
            inicio_prop=str(d.fecha_inicio or ""),
            inicio_contrato=str(d.contrato_inicio or ""),
            plaza_prop=d.id_plaza, plaza_contrato=d.contrato_plaza,
            clausula=d.clausula_efectiva, computa="sí" if d.computa else "no",
            incidencias=" · ".join(probl),
        ))
    _hoja(wb, "A12_Coherencia_enlace_directo", [
        ("propuesta_id", "Propuesta"),
        ("idrh_prop", "DNI propuesta"), ("idrh_contrato", "DNI contrato"),
        ("inicio_prop", "Inicio propuesta"), ("inicio_contrato", "Inicio contrato"),
        ("plaza_prop", "Plaza propuesta"), ("plaza_contrato", "Plaza contrato"),
        ("clausula", "Cláusula"), ("computa", "¿Computa?"),
        ("incidencias", "Incidencias"),
    ], incoh)

    # A13 Control: 'Días Usados' de PeopleNet vs días comprometidos por contrato
    if res.usados_vs_contratos:
        _hoja(wb, "A13_Usados_vs_Contratos", [
            ("clausula", "Cláusula"),
            ("comprometido_contratos", "Comprometido a fin real (contratos)"),
            ("dias_usados_peoplenet", "Días Usados (PeopleNet)"),
            ("diferencia", "Diferencia"),
            ("coincidencia_pct", "Coincidencia %"),
            ("estado", "Estado"),
        ], res.usados_vs_contratos)

    # B1/B2 Modo ANCLADO A PEOPLENET (disponible = Contratación − Usados_real −
    # Pendiente). Control por DIRECCIÓN (verde si su global ≥ 0).
    if res.anclado_direccion:
        _hoja(wb, "B1_Anclado_PeopleNet_Direccion", [
            ("clausula", "Cláusula"), ("direccion", "Dir."),
            ("direccion_nombre", "Dirección"),
            ("contratacion", "Contratación (reparto)"),
            ("usados_real", "Usados real (PeopleNet)"),
            ("pendiente", "Pendiente (reservas)"),
            ("disponible", "Disponible"), ("estado", "Estado"),
        ], res.anclado_direccion)
    if res.anclado_plaza:
        _hoja(wb, "B2_Anclado_PeopleNet_Plaza", [
            ("clausula", "Cláusula"), ("direccion", "Dir."),
            ("id_plaza", "ID Plaza"),
            ("contratacion_plaza", "Contratación (reparto)"),
            ("usados_real", "Usados real (PeopleNet)"),
            ("pendiente", "Pendiente (reservas)"),
            ("disponible", "Disponible"),
        ], res.anclado_plaza)
    if res.anclado_sin_gfh:
        _hoja(wb, "B3_Sin_GFH", [
            ("idrh", "DNI"), ("id_plaza", "ID Plaza"), ("clausula", "Cláusula"),
            ("num_periodo", "Nº periodo"), ("inicio", "Inicio"), ("fin", "Fin"),
            ("dias", "Días sin división"), ("gfh", "GFH (falta en el maestro)"),
        ], res.anclado_sin_gfh)
    if res.saldo_ini_direccion:
        _hoja(wb, "C1_Saldo_01_01_Direccion", _COLS_SALDO_INI_DIR,
              res.saldo_ini_direccion)
    if res.saldo_ini_plaza:
        _hoja(wb, "C2_Saldo_01_01_Plaza", _COLS_SALDO_INI_PLAZA,
              res.saldo_ini_plaza)

    Path(ruta).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return Path(ruta)


def exporta_recarga(
    res: ResultadoConciliacion, ruta: str | Path
) -> Path:
    """Fichero de recarga completa de Propuestas (incluye filas con saldo 0)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Recarga"
    columnas = [
        ("id_plaza", "ID Plaza"), ("descripcion_categoria", "Descripción categoría"),
        ("direccion_codigo", "Dirección"), ("clausula", "Cláusula"),
        ("saldo_calculado", "Saldo para cargar"),
    ]
    for j, (_, tit) in enumerate(columnas, start=1):
        c = ws.cell(row=1, column=j, value=tit)
        c.font = _CAB
        c.fill = _FILL_CAB
    for i, f in enumerate(res.detalle, start=2):
        vals = [f.id_plaza, f.descripcion_categoria, f.direccion_codigo,
                f.clausula, f.saldo_calculado]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v)
    ws.freeze_panes = "A2"
    Path(ruta).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return Path(ruta)


def exporta_recarga_anclada(
    res: ResultadoConciliacion, ruta: str | Path
) -> Path:
    """Recarga en modo ANCLADO A PEOPLENET: disponible por plaza =
    Contratación (reparto del corte) − Usados_real − Pendiente."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Recarga_anclada"
    columnas = [
        ("id_plaza", "ID Plaza"), ("direccion", "Dirección"),
        ("clausula", "Cláusula"), ("disponible", "Saldo para cargar"),
    ]
    for j, (_, tit) in enumerate(columnas, start=1):
        c = ws.cell(row=1, column=j, value=tit)
        c.font = _CAB
        c.fill = _FILL_CAB
    for i, f in enumerate(res.anclado_plaza, start=2):
        vals = [f["id_plaza"], f["direccion"], f["clausula"], f["disponible"]]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v)
    ws.freeze_panes = "A2"
    Path(ruta).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return Path(ruta)


def exporta_saldo_inicial(res: ResultadoConciliacion, ruta: str | Path) -> Path:
    """Hoja de cálculo del saldo inicial 01/01/2026 reconstruido (anclado a
    PeopleNet): pestañas C0 método, C1 por dirección, C2 por plaza."""
    wb = Workbook()
    wb.remove(wb.active)

    ws0 = wb.create_sheet("C0_Método")
    metodo = [
        ("Saldo inicial 01/01/2026 anclado a PeopleNet", True),
        ("", False),
        ("Reconstruye, por plaza y dirección, el saldo a 01/01/2026 de forma que", False),
        ("en el corte (02/07/2026) coincida EXACTAMENTE con el saldo real", False),
        ("(postcontrol), y proyecta el saldo actual solo con movimientos de", False),
        ("PeopleNet (contratos), olvidando las propuestas comprometidas.", False),
        ("", False),
        ("Fórmulas por (plaza, dirección, cláusula):", True),
        ("  Saldo inicial 01/01  = Saldo en el corte + Consumo PeopleNet al corte", False),
        ("  Saldo en el corte    = postcontrol del 02/07 (coincide por construcción)", False),
        ("  Saldo actual solo-PN = Saldo en el corte − Consumo PeopleNet desde el corte", False),
        ("                       = Saldo inicial 01/01 − Consumo PeopleNet total", False),
        ("", False),
        ("Consumo PeopleNet = días de cada contrato a su fin real (tope 31/12),", False),
        ("igual que 'Días Usados' de PeopleNet, repartidos a la división real por", False),
        ("tramo GFH. El corte de un contrato entre 'al corte' y 'desde el corte' se", False),
        ("hace por su FECHA DE INICIO (≤ 02/07 = al corte; posterior = desde el", False),
        ("corte), robusta; no por la fecha de alta reconstruida.", False),
        ("", False),
        ("NO incluye las propuestas comprometidas sin contrato (reservas): es la", False),
        ("foto solo-PeopleNet. El disponible con reservas está en el modo anclado", False),
        ("(pestañas B1–B3 de la conciliación).", False),
    ]
    for txt, negrita in metodo:
        ws0.append([txt])
        if negrita:
            ws0.cell(ws0.max_row, 1).font = Font(bold=True, size=12)
    ws0.column_dimensions["A"].width = 92

    ws1 = _hoja(wb, "C1_Saldo_01_01_Direccion", _COLS_SALDO_INI_DIR,
                res.saldo_ini_direccion)
    ws2 = _hoja(wb, "C2_Saldo_01_01_Plaza", _COLS_SALDO_INI_PLAZA,
                res.saldo_ini_plaza)
    # color del saldo actual (última columna)
    for ws, cols in ((ws1, _COLS_SALDO_INI_DIR), (ws2, _COLS_SALDO_INI_PLAZA)):
        col_act = len(cols)
        for i in range(2, ws.max_row + 1):
            v = ws.cell(i, col_act).value
            if isinstance(v, (int, float)):
                ws.cell(i, col_act).fill = _COLOR["ROJO" if v < 0 else "VERDE"]
    # totales por cláusula al pie de C1
    ws1.append([])
    for cl in sorted({r["clausula"] for r in res.saldo_ini_direccion}):
        filas_cl = [r for r in res.saldo_ini_direccion if r["clausula"] == cl]
        ws1.append([f"TOTAL {cl}", "", "",
                    sum(r["saldo_01_01"] for r in filas_cl),
                    sum(r["usados_corte"] for r in filas_cl),
                    sum(r["saldo_corte"] for r in filas_cl),
                    sum(r["usados_desde"] for r in filas_cl),
                    sum(r["saldo_actual"] for r in filas_cl)])
        for c in ws1[ws1.max_row]:
            c.font = Font(bold=True)

    Path(ruta).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return Path(ruta)
